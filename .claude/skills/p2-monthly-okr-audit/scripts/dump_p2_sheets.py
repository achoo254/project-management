"""Dump P2 evaluation + rubric sheets of a member file to a UTF-8 txt.

Usage:
    python dump_p2_sheets.py "<path to P2 .xlsx>" "<out.txt>"

Reads (data_only=cached values):
  - "ĐÁNH GIÁ HIỆU SUẤT P2"  -> all non-empty cells (3 Delivery levels at C17/C18/C19,
                               totals C20/C23/C24, metadata B3/B4/C3-D3/C4-D4)
  - "CẤU TRÚC & TRỌNG SỐ"     -> rubric (role weights, level->score, delivery weights, thresholds)
Also dumps formulas of the eval sheet so the scorer formula can be verified.

Note: if the file is open in Excel (~$lock present), openpyxl reads the last
saved version on disk (may be slightly stale) — that is acceptable for audit.
"""
import sys
from openpyxl import load_workbook

if len(sys.argv) < 3:
    print("Usage: dump_p2_sheets.py <xlsx_path> <out_txt>")
    sys.exit(1)

path, out_path = sys.argv[1], sys.argv[2]
out = open(out_path, "w", encoding="utf-8")
def p(*a): print(*a, file=out)

wbv = load_workbook(path, data_only=True)
wbf = load_workbook(path, data_only=False)
p("=== SHEETS ===", wbv.sheetnames)

EVAL = "ĐÁNH GIÁ HIỆU SUẤT P2"
RUBRIC = "CẤU TRÚC & TRỌNG SỐ"

def dump_cells(sheet_name, title):
    if sheet_name not in wbv.sheetnames:
        p(f"\n### MISSING SHEET {sheet_name!r}")
        return
    ws = wbv[sheet_name]
    p(f"\n############ {title}: {sheet_name!r}  dims={ws.dimensions} ############")
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is not None and str(v).strip() != "":
                p(f"{cell.coordinate}\t{repr(v)}")

# eval sheet (fuzzy match if exact name differs)
eval_name = EVAL if EVAL in wbv.sheetnames else next(
    (s for s in wbv.sheetnames if "HIỆU SUẤT" in s.upper() and "P2" in s.upper()), None)
dump_cells(eval_name, "EVAL")
dump_cells(RUBRIC, "RUBRIC")

# formulas of eval sheet (verify scoring)
if eval_name:
    p(f"\n############ FORMULAS in {eval_name!r} ############")
    wf = wbf[eval_name]
    for row in wf.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                p(f"{cell.coordinate}\t{v}")

out.close()
print(f"OK -> {out_path}")
