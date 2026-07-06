"""Dump P2 evaluation sheet + key metadata for any member file. Usage: dump_p2_member.py <xlsx_path> <out_txt>"""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
out_path = sys.argv[2]
out = open(out_path, "w", encoding="utf-8")
def p(*a):
    print(*a, file=out)

wbv = load_workbook(path, data_only=True)
p("=== SHEETS ===", wbv.sheetnames)

sheet = "ĐÁNH GIÁ HIỆU SUẤT P2"
if sheet not in wbv.sheetnames:
    sheet = next((s for s in wbv.sheetnames if "P2" in s.upper() and "HIỆU SUẤT" in s.upper()), wbv.sheetnames[0])

ws = wbv[sheet]
p(f"\n=== {sheet!r} dims={ws.dimensions} ===")
for row in ws.iter_rows():
    for cell in row:
        v = cell.value
        if v is not None and str(v).strip() != "":
            p(f"{cell.coordinate}\t{repr(v)}")
