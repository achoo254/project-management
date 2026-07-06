"""Dump rubric + appendix sheets of Dat's P2 to verify scoring formula & transparency."""
from openpyxl import load_workbook

out = open(r"D:\CONG VIEC\project-management\scripts\p2_dat_all_dump.txt", "w", encoding="utf-8")
def p(*a):
    print(*a, file=out)

path = r"D:\CONG VIEC\project-management\members\T5\P2_Form_BP_Cong_Nghe - Dat.xlsx"

# data_only=True -> cached values; also load formulas to see how totals computed
wb_val = load_workbook(path, data_only=True)
wb_fml = load_workbook(path, data_only=False)

for sheet in ["CẤU TRÚC & TRỌNG SỐ", "MÔ TẢ - PHỤ LỤC", "Summary_Ranking", "TỔNG QUAN & LƯU Ý"]:
    if sheet not in wb_val.sheetnames:
        p(f"\n### MISSING SHEET {sheet!r}")
        continue
    wsv = wb_val[sheet]
    p(f"\n############ SHEET: {sheet!r}  dims={wsv.dimensions} ############")
    for row in wsv.iter_rows():
        for cell in row:
            v = cell.value
            if v is not None and str(v).strip() != "":
                p(f"{cell.coordinate}\t{repr(v)}")

# Formulas for the P2 evaluation sheet (to verify totals)
p("\n############ FORMULAS in 'ĐÁNH GIÁ HIỆU SUẤT P2' ############")
wsf = wb_fml["ĐÁNH GIÁ HIỆU SUẤT P2"]
for row in wsf.iter_rows():
    for cell in row:
        v = cell.value
        if isinstance(v, str) and v.startswith("="):
            p(f"{cell.coordinate}\t{v}")
