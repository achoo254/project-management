# -*- coding: utf-8 -*-
"""Sửa P2+ Tổng và xếp loại sau khi đã chèn dòng 15: tìm dòng P2+ Tổng và đặt công thức B=(D11+B31)/2, C=ngưỡng."""
from __future__ import annotations

import os
import sys

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(BASE, "members", "P2_Form_BP_Cong_Nghe - Bui Van Ho.xlsx")

def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    wb = openpyxl.load_workbook(PATH, data_only=False)
    ws = wb["Product_Engineer"]
    # Tìm dòng có "P2+ Tổng" trong cột A
    p2_total_row = None
    for r in range(28, 36):
        if ws.cell(r, 1).value and "P2+" in str(ws.cell(r, 1).value) and "Tổng" in str(ws.cell(r, 1).value):
            p2_total_row = r
            break
    if p2_total_row is None:
        print("Không tìm thấy dòng P2+ Tổng.")
        wb.close()
        return
    # B = (D11 + B31) / 2  (P2-Delivery đang ở B31 sau khi chèn)
    ws.cell(p2_total_row, 2).value = '=IF(OR(D11="",B31=""),"",(D11+B31)/2)'
    # C = ngưỡng theo giá trị ô B cùng dòng
    cell_b = f"B{p2_total_row}"
    for m in list(ws.merged_cells.ranges):
        if m.min_row <= p2_total_row <= m.max_row and m.min_col <= 3 <= m.max_col:
            ws.unmerge_cells(str(m))
            break
    ws.cell(p2_total_row, 3).value = f'=IF({cell_b}="","",IF({cell_b}<1,"Chưa đạt",IF({cell_b}<=1,"Đạt",IF({cell_b}<=1.2,"Tốt",""))))'
    wb.save(PATH)
    wb.close()
    print(f"Đã đặt P2+ Tổng và xếp loại tại dòng {p2_total_row}.")

if __name__ == "__main__":
    main()
