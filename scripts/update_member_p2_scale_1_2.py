# -*- coding: utf-8 -*-
"""
Cập nhật P2+ trong file member (Bui Van Ho) theo thang 1.2:
- P2-Core tối đa 1.2: L1=0.36, L2=0.84, L3=1.2 (trọng số % giữ nguyên).
- P2-Delivery tối đa 1.2: Chưa đạt=0, Đạt=24/35 (~0.686), Tốt=1.2.
- P2+ Tổng = (P2-Core + P2-Delivery) / 2 → max 1.2.
- 3 mốc: <0.9 Chưa đạt, 0.9–1.0 Đạt, >1.0–1.2 Tốt (hiển thị ở C31).
"""
from __future__ import annotations

import sys

import openpyxl

MEMBER_PATH = r"c:\CONG VIEC\project-management\members\P2_Form_BP_Cong_Nghe - Bui Van Ho.xlsx"
SHEET = "Product_Engineer"

# P2-Core: thang 1.2 (L1=0.36, L2=0.84, L3=1.2)
CORE_L1, CORE_L2, CORE_L3 = 0.36, 0.84, 1.2

# P2-Delivery: Chưa đạt=0, Đạt=24/35, Tốt=1.2 (map từ 1.5/3.5/5 → 0 / 0.686 / 1.2)
DELIVERY_CHUA_DAT = 0
DELIVERY_DAT = "24/35"  # trong công thức Excel
DELIVERY_TOT = 1.2


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    wb = openpyxl.load_workbook(MEMBER_PATH, data_only=False)
    if SHEET not in wb.sheetnames:
        print("Sheet Product_Engineer không tồn tại.")
        wb.close()
        return

    ws = wb[SHEET]

    # 1. D7:D10 – Điểm quy đổi P2-Core (thang 1.2)
    for r in range(7, 11):
        ws.cell(r, 4).value = (
            f'=IF(C{r}="","",B{r}*IF(C{r}="L1",{CORE_L1},IF(C{r}="L2",{CORE_L2},IF(C{r}="L3",{CORE_L3},0)))/100)'
        )

    # 2. D11 giữ nguyên = SUM(D7:D10) → đã là tổng thang 1.2
    ws.cell(11, 4).value = "=SUM(D7:D10)"

    # 3. B30 – P2-Delivery thang 1.2 (Chưa đạt=0, Đạt=24/35, Tốt=1.2)
    part25 = f'IF(C25="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C25="Đạt",{DELIVERY_DAT},IF(C25="Tốt",{DELIVERY_TOT},0)))'
    part26 = f'IF(C26="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C26="Đạt",{DELIVERY_DAT},IF(C26="Tốt",{DELIVERY_TOT},0)))'
    part27 = f'IF(C27="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C27="Đạt",{DELIVERY_DAT},IF(C27="Tốt",{DELIVERY_TOT},0)))'
    count_part = "(C25<>\"\")+(C26<>\"\")+(C27<>\"\")"
    ws.cell(30, 2).value = (
        f'=IF({count_part}=0,"",(({part25})+({part26})+({part27}))/({count_part}))'
    )

    # 4. B31 – P2+ Tổng = (D11 + B30) / 2 (cả hai đã 0–1.2)
    ws.cell(31, 2).value = '=IF(OR(D11="",B30=""),"",(D11+B30)/2)'

    # 5. Xếp mức 3 mốc: <0.9 Chưa đạt, 0.9–1.0 Đạt, >1.0 Tốt
    # C31 có thể nằm trong merged range → bỏ merge rồi ghi công thức vào C31
    from openpyxl.utils import get_column_letter
    to_unmerge = []
    for m in list(ws.merged_cells.ranges):
        if m.min_row <= 31 <= m.max_row and m.min_col <= 3 <= m.max_col:
            to_unmerge.append(str(m))
    for ref in to_unmerge:
        ws.unmerge_cells(ref)
    ws.cell(31, 3).value = '=IF(B31="","",IF(B31<0.9,"Chưa đạt",IF(B31<=1,"Đạt","Tốt")))'

    wb.save(MEMBER_PATH)
    wb.close()
    print("Đã cập nhật: P2-Core & P2-Delivery max 1.2, P2+ Tổng = (Core+Delivery)/2, 3 mốc Chưa đạt/Đạt/Tốt tại C31.")


if __name__ == "__main__":
    main()
