# -*- coding: utf-8 -*-
"""Kiểm tra lỗi file Bui Van Ho: công thức, tham chiếu, merge, data validation."""
import sys
import os
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import openpyxl

base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
path = os.path.join(base, "members", "P2_Form_BP_Cong_Nghe - Bui Van Ho.xlsx")

print("File:", path)
print("Exists:", os.path.exists(path))
if not os.path.exists(path):
    sys.exit(1)

wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["Product_Engineer"]

print("\n--- 1. Công thức D7:D11, B30-B32, C31-C32 ---")
for r, c in [(7,4),(8,4),(9,4),(10,4),(11,4),(30,2),(31,2),(31,3),(32,2),(32,3)]:
    cell = ws.cell(r, c)
    v = cell.value
    if getattr(cell, 'data_type', None) == 'f':
        print(f"  Row{r} Col{c}: [FORMULA] {v}")
    else:
        print(f"  Row{r} Col{c}: {v!r}")

print("\n--- 2. Nội dung quanh C25:C28 (Delivery) ---")
for r in range(23, 32):
    a, b, c = ws.cell(r,1).value, ws.cell(r,2).value, ws.cell(r,3).value
    print(f"  {r}: A={str(a)[:40]!r}  B={str(b)[:30]!r}  C={str(c)[:40]!r}")

print("\n--- 3. Merged cells ---")
for m in list(ws.merged_cells.ranges):
    print(f"  {m}")

print("\n--- 4. Data validations ---")
for dv in ws.data_validations.dataValidation:
    print(f"  sqref={dv.sqref} formula1={dv.formula1!r}")

print("\n--- 5. Giá trị tính ra (data_only=True) ---")
wb2 = openpyxl.load_workbook(path, data_only=True)
ws2 = wb2["Product_Engineer"]
for r, c in [(7,4),(8,4),(9,4),(10,4),(11,4),(30,2),(31,2),(31,3),(32,2),(32,3)]:
    v = ws2.cell(r, c).value
    if v is not None and isinstance(v, str) and (v.startswith("#") or "ERROR" in v.upper()):
        print(f"  [ERROR] Row{r} Col{c}: {v}")
    else:
        print(f"  Row{r} Col{c} (computed): {v!r}")
wb2.close()

wb.close()
print("\nDone.")
