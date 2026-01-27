import sys
import openpyxl


def shorten(v: object, n: int = 40) -> str:
    if v is None:
        return ""
    s = str(v).replace("\n", " ").strip()
    return s if len(s) <= n else s[: n - 1] + "…"


def main() -> None:
    # Avoid UnicodeEncodeError on Windows consoles with legacy codepages.
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    path = r"c:\CONG VIEC\project-management\P2_Form_Professional_Report_JuniorPlus_Table.xlsx"
    wb = openpyxl.load_workbook(path)
    print("sheets:", wb.sheetnames)

    for name in ["Product_Engineer", "PE_Criteria_Detail", "PE_Criteria_Table", "Summary_Ranking"]:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        merges = list(ws.merged_cells.ranges)
        tables = []
        try:
            tables = list(getattr(ws, "tables", {}).keys())
        except Exception:
            pass

        print(f"\n== {name} ==")
        print("dimensions:", ws.dimensions)
        print("freeze_panes:", ws.freeze_panes)
        print("max_row:", ws.max_row, "max_col:", ws.max_column)
        print("merged_ranges:", len(merges))
        print("tables:", tables)
        if merges:
            print("first_merges:", [str(m) for m in merges[:10]])

        shown = 0
        for r in range(1, min(ws.max_row, 300) + 1):
            vals = [shorten(ws.cell(r, c).value, 50) for c in range(1, 13)]
            if any(vals):
                print(r, vals)
                shown += 1
                if shown >= 35:
                    break

    wb.close()


if __name__ == "__main__":
    main()

