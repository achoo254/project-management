# -*- coding: utf-8 -*-
"""
Áp dụng toàn bộ cập nhật giống file Bui Van Ho cho: Dinh Tien Chuong, Nguyen Dang Duc, Nguyen Huy Hoang.
- L0 + THANG (L0/L1/L2/L3), Mức C7:C10, D7:D10 công thức, Điểm text C26:C28
- P2-Delivery từ C26,C27,C28 (0.8/1/1.2), validation C26:C28
- P2+ Tổng = (D11+B31)/2, xếp loại <1/1/1.2
- Sheet Cấu trúc & Trọng số: 3 dòng ngưỡng
"""
from __future__ import annotations

import os
import sys

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MEMBERS_DIR = os.path.join(BASE_DIR, "members")

SHEET_PE = "Product_Engineer"
SHEET_CT = "Cấu trúc & Trọng số"

# File cần cập nhật (không có Bui Van Ho vì đã xong)
MEMBER_FILES = [
    "P2_Form_BP_Cong_Nghe - Dinh Tien Chuong.xlsx",
    "P2_Form_BP_Cong_Nghe - Nguyen Dang Duc.xlsx",
    "P2_Form_BP_Cong_Nghe - Nguyen Huy Hoang.xlsx",
]

CORE_L0, CORE_L1, CORE_L2, CORE_L3 = 0.8, 0.9, 1, 1.2
DELIVERY_CHUA_DAT, DELIVERY_DAT, DELIVERY_TOT = 0.8, 1, 1.2


def apply_to_file(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    if SHEET_PE not in wb.sheetnames:
        wb.close()
        return
    ws = wb[SHEET_PE]

    # Bỏ merge các vùng có thể ghi (tránh lỗi MergedCell read-only)
    for m in list(ws.merged_cells.ranges):
        try:
            ws.unmerge_cells(str(m))
        except Exception:
            pass

    # --- Đã có L0 chưa? (row 15 có "L0" thì coi như đã insert)
    already_has_l0 = ws.cell(15, 1).value and str(ws.cell(15, 1).value).strip() == "L0"

    # 1. Header Mức (row 6)
    if ws.cell(6, 3).value:
        ws.cell(6, 3).value = "Mức (L0/L1/L2/L3)"

    # 2. THANG + chèn L0 (chỉ chèn nếu chưa có)
    if ws.cell(13, 1).value and "THANG" in str(ws.cell(13, 1).value):
        ws.cell(13, 1).value = "THANG ĐÁNH GIÁ CHI TIẾT (L0/L1/L2/L3)"
    if not already_has_l0:
        ws.insert_rows(15)
        ws.cell(15, 1).value = "L0"
        ws.cell(15, 2).value = "Chưa đạt kết quả hoặc năng lực chưa đạt."

    # 3. Mô tả Điểm (tìm ô có "Điểm" và "Chưa đạt" trong cột C)
    for r in range(23, 29):
        val = ws.cell(r, 3).value
        if val and "Điểm" in str(val) and "Chưa đạt" in str(val):
            ws.cell(r, 3).value = "Điểm (Chưa đạt < 1, Đạt >=1, Tốt=1.2)"
            break

    # 4. Data validation C7:C10 = L0,L1,L2,L3
    to_remove = [dv for dv in list(ws.data_validations.dataValidation) if dv.sqref and ("C7:C10" in str(dv.sqref) or "C7" in str(dv.sqref))]
    for dv in to_remove:
        ws.data_validations.dataValidation.remove(dv)
    dv_core = DataValidation(type="list", formula1='"L0,L1,L2,L3"', allow_blank=True)
    dv_core.add("C7:C10")
    ws.add_data_validation(dv_core)

    # 5. D7:D10, D11
    for r in range(7, 11):
        ws.cell(r, 4).value = (
            f'=IF(C{r}="","",B{r}*IF(C{r}="L0",{CORE_L0},IF(C{r}="L1",{CORE_L1},IF(C{r}="L2",{CORE_L2},IF(C{r}="L3",{CORE_L3},0))))/100)'
        )
    ws.cell(11, 4).value = "=SUM(D7:D10)"

    # 6. P2-Delivery: tìm dòng có "P2-Delivery" trong A, ghi công thức C26,C27,C28 vào cột B
    delivery_row = None
    for r in range(28, 35):
        a = ws.cell(r, 1).value
        if a and "P2-Delivery" in str(a) and "Trung bình" not in str(a):
            delivery_row = r
            break
    if delivery_row is None:
        for r in range(28, 35):
            if ws.cell(r, 1).value and str(ws.cell(r, 1).value).strip() == "P2-Delivery":
                delivery_row = r
                break
    if delivery_row is not None:
        for m in list(ws.merged_cells.ranges):
            if m.min_row <= delivery_row <= m.max_row and m.min_col <= 2 <= m.max_col:
                ws.unmerge_cells(str(m))
                break
        part26 = f'IF(C26="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C26="Đạt",{DELIVERY_DAT},IF(C26="Tốt",{DELIVERY_TOT},0)))'
        part27 = f'IF(C27="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C27="Đạt",{DELIVERY_DAT},IF(C27="Tốt",{DELIVERY_TOT},0)))'
        part28 = f'IF(C28="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C28="Đạt",{DELIVERY_DAT},IF(C28="Tốt",{DELIVERY_TOT},0)))'
        count_part = '(C26<>"")+(C27<>"")+(C28<>"")'
        ws.cell(delivery_row, 2).value = f'=IF({count_part}=0,"",(({part26})+({part27})+({part28}))/({count_part}))'

    # 7. Data validation C26:C28 = Chưa đạt,Đạt,Tốt (bỏ C25:C27 cũ nếu có)
    to_remove_d = [dv for dv in list(ws.data_validations.dataValidation) if dv.sqref and "C25:C27" in str(dv.sqref)]
    for dv in to_remove_d:
        ws.data_validations.dataValidation.remove(dv)
    dv_del = DataValidation(type="list", formula1='"Chưa đạt,Đạt,Tốt"', allow_blank=True)
    dv_del.add("C26:C28")
    ws.add_data_validation(dv_del)

    # 8. P2+ Tổng: tìm dòng "P2+ Tổng", B = (D11+B31)/2, C = ngưỡng (B31 là P2-Delivery sau khi chèn)
    p2_total_row = None
    for r in range(28, 36):
        if ws.cell(r, 1).value and "P2+" in str(ws.cell(r, 1).value) and "Tổng" in str(ws.cell(r, 1).value):
            p2_total_row = r
            break
    if p2_total_row is not None:
        for m in list(ws.merged_cells.ranges):
            if m.min_row <= p2_total_row <= m.max_row and (m.min_col <= 2 <= m.max_col or m.min_col <= 3 <= m.max_col):
                ws.unmerge_cells(str(m))
        ws.cell(p2_total_row, 2).value = '=IF(OR(D11="",B31=""),"",(D11+B31)/2)'
        cell_b = f"B{p2_total_row}"
        ws.cell(p2_total_row, 3).value = f'=IF({cell_b}="","",IF({cell_b}<1,"Chưa đạt",IF({cell_b}<=1,"Đạt",IF({cell_b}<=1.2,"Tốt",""))))'

    # 9. Sheet Cấu trúc & Trọng số
    if SHEET_CT in wb.sheetnames:
        ws_ct = wb[SHEET_CT]
        header_row = None
        for r in range(1, 50):
            val = ws_ct.cell(r, 1).value
            if val and "NGƯỠNG XẾP LOẠI" in str(val).upper():
                header_row = r + 1
                break
        if header_row is None:
            for r in range(1, 50):
                if ws_ct.cell(r, 1).value and str(ws_ct.cell(r, 1).value).strip() == "P2+ Tổng":
                    header_row = r
                    break
        if header_row is not None:
            ws_ct.cell(header_row, 1).value = "P2+ Tổng"
            ws_ct.cell(header_row, 2).value = "Xếp loại"
            ws_ct.cell(header_row + 1, 1).value = "< 1"
            ws_ct.cell(header_row + 1, 2).value = "Chưa đạt"
            ws_ct.cell(header_row + 2, 1).value = "1"
            ws_ct.cell(header_row + 2, 2).value = "Đạt"
            ws_ct.cell(header_row + 3, 1).value = "> 1 và ≤ 1.2"
            ws_ct.cell(header_row + 3, 2).value = "Tốt"
            for c in (1, 2):
                if ws_ct.cell(header_row + 4, c).value is not None:
                    ws_ct.cell(header_row + 4, c).value = None

    wb.save(path)
    wb.close()


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    for name in MEMBER_FILES:
        path = os.path.join(MEMBERS_DIR, name)
        if not os.path.exists(path):
            print(f"Bỏ qua (không tìm thấy): {name}")
            continue
        try:
            apply_to_file(path)
            print(f"Đã cập nhật: {name}")
        except Exception as e:
            print(f"Lỗi {name}: {e}")
    print("Xong.")


if __name__ == "__main__":
    main()
