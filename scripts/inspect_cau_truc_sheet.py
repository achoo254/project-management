# -*- coding: utf-8 -*-
import sys
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass
import openpyxl
path = r"c:\CONG VIEC\project-management\members\P2_Form_BP_Cong_Nghe - Bui Van Ho.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
if "Cấu trúc & Trọng số" in wb.sheetnames:
    ws = wb["Cấu trúc & Trọng số"]
    for r in range(1, 50):
        row = [ws.cell(r, c).value for c in range(1, 6)]
        if any(x is not None and str(x).strip() for x in row):
            print(r, row)
wb.close()
