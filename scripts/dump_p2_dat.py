"""Dump P2 evaluation sheet for Dat to inspect P2 line items vs OKR/KPI."""
import sys
from openpyxl import load_workbook

out = open(r"D:\CONG VIEC\project-management\scripts\p2_dat_dump.txt", "w", encoding="utf-8")
def p(*a):
    print(*a, file=out)

path = r"D:\CONG VIEC\project-management\members\T5\P2_Form_BP_Cong_Nghe - Dat.xlsx"
wb = load_workbook(path, data_only=True)

p("=== SHEETS ===")
for s in wb.sheetnames:
    p(repr(s))

target = "ĐÁNH GIÁ HIỆU SUẤT P2"
# Find best matching sheet name
match = None
for s in wb.sheetnames:
    if s.strip().upper().startswith("ĐÁNH GIÁ HIỆU SUẤT") or "HIỆU SUẤT P2" in s.upper():
        match = s
        break
if match is None:
    # fallback: any sheet containing P2
    for s in wb.sheetnames:
        if "P2" in s.upper():
            match = s
            break

p(f"\n=== DUMP SHEET: {match!r} ===")
ws = wb[match]
p(f"dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
p("--- merged cells ---")
for mc in ws.merged_cells.ranges:
    p(mc)
p("--- non-empty cells ---")
for row in ws.iter_rows():
    for cell in row:
        v = cell.value
        if v is not None and str(v).strip() != "":
            p(f"{cell.coordinate}\t{repr(v)}")
