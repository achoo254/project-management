# -*- coding: utf-8 -*-
"""
Áp dụng 3 luật cap P2-Delivery (chặn chạy số) trong sheet Product_Engineer:
- Luật 1: Quality (C26) = Chưa đạt → Delivery tối đa 2
- Luật 2: OKR (C25) = Chưa đạt → Delivery tối đa 2
- Luật 3: Sự cố P0 do chủ quan → chọn cap 1 hoặc 2 tại C28

Thêm hàng 28 (nhãn + C28 dropdown), ghi chú 3 luật, công thức B30 = MIN(raw_avg, cap).
"""
from __future__ import annotations

import sys
import os
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = r"c:\CONG VIEC\project-management"
PATH_MAIN = os.path.join(ROOT, "P2_Form_BP_Cong_Nghe.xlsx")
MEMBERS_DIR = os.path.join(ROOT, "members")
MEMBER_FILES = [
    "P2_Form_BP_Cong_Nghe - Bui Van Ho.xlsx",
    "P2_Form_BP_Cong_Nghe - Dinh Tien Chuong.xlsx",
    "P2_Form_BP_Cong_Nghe - Nguyen Dang Duc.xlsx",
    "P2_Form_BP_Cong_Nghe - Nguyen Huy Hoang.xlsx",
]
SHEET = "Product_Engineer"

# C28: Luật 3 — Sự cố P0 do chủ quan
C28_LIST = "P0 - cap 2,P0 - cap 1"  # allow_blank = True cho "Không áp dụng"


def apply_cap_rules(path: str) -> bool:
    """Thêm hàng 28 (C28), ghi chú 3 luật, cập nhật công thức B30 với cap."""
    wb = openpyxl.load_workbook(path, data_only=False)
    if SHEET not in wb.sheetnames:
        print(f"  Skip (no sheet '{SHEET}'): {path}")
        wb.close()
        return False

    ws = wb[SHEET]

    # 1. Hàng 28: nhãn (A28:B28 merge) + ô C28 + data validation
    ws.cell(28, 1).value = "Sự cố P0 do chủ quan (tùy chọn)"
    try:
        ws.merge_cells(start_row=28, start_column=1, end_row=28, end_column=2)
    except Exception:
        pass
    dv_p0 = DataValidation(
        type="list",
        formula1=f'"{C28_LIST}"',
        allow_blank=True,
    )
    dv_p0.add("C28")
    ws.add_data_validation(dv_p0)

    # 2. Ghi chú 3 luật: thêm vào nhãn hàng 28 (merge A28:B28) dòng ngắn
    ws.cell(28, 1).value = (
        "Sự cố P0 do chủ quan (tùy chọn). "
        "Luật 1: Quality Chưa đạt → cap 2. Luật 2: OKR Chưa đạt → cap 2. Luật 3: Chọn tại C28."
    )

    # 3. Công thức B30: raw_avg rồi MIN(raw_avg, cap)
    sum_part = (
        'IF(C25="",0,IF(C25="Chưa đạt",1,IF(C25="Đạt",2,IF(C25="Tốt",3,0))))+'
        'IF(C26="",0,IF(C26="Chưa đạt",1,IF(C26="Đạt",2,IF(C26="Tốt",3,0))))+'
        'IF(C27="",0,IF(C27="Chưa đạt",1,IF(C27="Đạt",2,IF(C27="Tốt",3,0))))'
    )
    count_part = '(C25<>"")+(C26<>"")+(C27<>"")'
    # cap: 1 nếu C28="P0 - cap 1"; 2 nếu C25/C26 Chưa đạt hoặc C28="P0 - cap 2"; else 3
    cap_part = 'IF(C28="P0 - cap 1",1,IF(OR(C25="Chưa đạt",C26="Chưa đạt",C28="P0 - cap 2"),2,3))'
    formula_b30 = (
        f'=IF({count_part}=0,"",MIN(({sum_part})/{count_part},{cap_part}))'
    )
    ws.cell(30, 2).value = formula_b30

    try:
        wb.save(path)
    except PermissionError:
        print(f"  Permission denied (file open?): {os.path.basename(path)}")
        wb.close()
        return False
    wb.close()
    return True


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    paths = [PATH_MAIN]
    for f in MEMBER_FILES:
        p = os.path.join(MEMBERS_DIR, f)
        if os.path.isfile(p):
            paths.append(p)

    for path in paths:
        name = os.path.basename(path)
        if apply_cap_rules(path):
            print(f"OK: {name}")
    print("Done: 3 luật cap Delivery đã áp dụng (hàng 28 + C28 + B30).")


if __name__ == "__main__":
    main()
