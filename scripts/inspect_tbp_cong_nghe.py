# -*- coding: utf-8 -*-
"""Inspect sheet TBP_CÔNG NGHỆ in P2_Form_BP_Cong_Nghe - LEAD.xlsx"""
import sys
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

import openpyxl

path = r"c:\CONG VIEC\project-management\P2_Form_BP_Cong_Nghe - LEAD.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
print("Sheets:", wb.sheetnames)

# Find TBP_CÔNG NGHỆ (may have different encoding)
sheet_name = None
for s in wb.sheetnames:
    if "TBP" in s and "CONG" in s.upper():
        sheet_name = s
        break
if not sheet_name:
    # try exact
    for s in wb.sheetnames:
        if "TBP" in s:
            sheet_name = s
            break
if not sheet_name:
    sheet_name = "TBP_CÔNG NGHỆ" if "TBP_CÔNG NGHỆ" in wb.sheetnames else wb.sheetnames[0]

ws = wb[sheet_name]
print("Using sheet:", sheet_name)
print("Dimensions:", ws.dimensions)
print()

# C7:C10 and context (headers, adjacent columns)
print("=== Rows 1-15, cols A-E (value and formula) ===")
for r in range(1, 16):
    row_vals = []
    for c in range(1, 6):
        cell = ws.cell(r, c)
        v = cell.value
        f = cell.value if isinstance(getattr(cell, 'data_type', None), str) and cell.data_type == 'f' else None
        # openpyxl: formula is in .value when data_type is 'f'
        if hasattr(cell, 'data_type') and cell.data_type == 'f':
            row_vals.append(f"f={v}")
        else:
            row_vals.append(str(v)[:50] if v is not None else "")
    print(r, " | ".join(row_vals))

print()
print("=== C7:C10 formulas (column C) ===")
for r in range(7, 11):
    cell = ws.cell(r, 3)
    print(f"  C{r}: value={cell.value!r} data_type={getattr(cell,'data_type',None)}")

# Also check if there's a "điểm quy đổi" column and its formula
print()
print("=== All columns in row 5-6 and 7-10 to find 'điểm quy đổi' ===")
for r in range(5, 11):
    vals = []
    for c in range(1, min(ws.max_column, 15) + 1):
        cell = ws.cell(r, c)
        v = cell.value
        if hasattr(cell, 'data_type') and cell.data_type == 'f':
            vals.append(f"[{c}]f:{v}")
        else:
            vals.append(f"[{c}]{v!r}")
    print("Row", r, ":", " | ".join(vals[:12]))

# Data validation on C7:C10
for r in range(7, 11):
    cell = ws.cell(r, 3)
    if cell.data_type == 'f':
        print("C%d formula: %s" % (r, cell.value))
    if hasattr(ws, 'data_validations') and ws.data_validations:
        for dv in ws.data_validations.dataValidation:
            print("DV: %s" % (dv))

# Rows 13-20 for scale description
print()
print("=== Rows 13-22 (scale description) ===")
for r in range(13, 23):
    row_vals = [str(ws.cell(r, c).value)[:60] for c in range(1, 4)]
    print(r, " | ".join(row_vals))

wb.close()
