from __future__ import annotations

from dataclasses import dataclass
from typing import Dict, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


@dataclass(frozen=True)
class CriterionRow:
    role: str
    category: str
    level: int
    description: str


def _safe_str(v: object) -> str:
    return "" if v is None else str(v).strip()


def _read_long_format(ws) -> List[CriterionRow]:
    rows: List[CriterionRow] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not any(r[:4]):
            continue
        role = _safe_str(r[0])
        category = _safe_str(r[1])
        try:
            level = int(r[2])
        except Exception:
            continue
        desc = _safe_str(r[3])
        if role and category and desc:
            rows.append(CriterionRow(role=role, category=category, level=level, description=desc))
    return rows


def _read_from_template() -> List[CriterionRow]:
    """
    Recover criteria content from template workbook.

    The template stores levels 1-5; the JuniorPlus table uses 1-4 which
    correspond to template levels 2-5 (shift -1).
    """
    template_path = r"c:\CONG VIEC\project-management\P2_Template_BP_Cong_Nghe_v2.xlsx"
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Tiêu chí"] if "Tiêu chí" in wb.sheetnames else wb[wb.sheetnames[2]]

    role_map = {"BE": "Backend", "FE": "FrontEnd", "DE": "Designer", "FS": "FullStack"}

    out: List[CriterionRow] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not any(r[:4]):
            continue
        role_code = _safe_str(r[0])
        category = _safe_str(r[1])
        try:
            lvl = int(r[2])
        except Exception:
            continue
        desc = _safe_str(r[3])

        # Map template levels 2..5 -> 1..4 for JuniorPlus table.
        if lvl < 2 or lvl > 5:
            continue
        lvl = lvl - 1

        role = role_map.get(role_code, role_code)
        if role and category and desc:
            out.append(CriterionRow(role=role, category=category, level=lvl, description=desc))

    wb.close()
    return out


def _ordered_unique(values: List[str]) -> List[str]:
    seen = set()
    out: List[str] = []
    for v in values:
        if v in seen:
            continue
        seen.add(v)
        out.append(v)
    return out


def _build_wide_matrix(
    rows: List[CriterionRow],
) -> Tuple[List[str], List[str], List[int], Dict[Tuple[str, str, int], str]]:
    roles = _ordered_unique([r.role for r in rows])
    categories = _ordered_unique([r.category for r in rows])
    levels = sorted({r.level for r in rows})
    cell: Dict[Tuple[str, str, int], str] = {}
    for r in rows:
        cell[(r.role, r.category, r.level)] = r.description
    return roles, categories, levels, cell


def _apply_border(ws, cell_range: str, border: Border) -> None:
    for row in ws[cell_range]:
        for c in row:
            c.border = border


def _set_col_widths(ws, widths: Dict[int, float]) -> None:
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def main() -> None:
    path = r"c:\CONG VIEC\project-management\P2_Form_Professional_Report_JuniorPlus_Table.xlsx"
    wb = openpyxl.load_workbook(path)

    old = wb["PE_Criteria_Detail"]
    old_idx = wb.sheetnames.index(old.title)
    long_rows = _read_long_format(old)
    if not long_rows:
        long_rows = _read_from_template()
    roles, categories, levels, cell = _build_wide_matrix(long_rows)

    # Recreate sheet to avoid leftover formatting.
    wb.remove(old)
    ws = wb.create_sheet("PE_Criteria_Detail", old_idx)

    # Colors / styles
    fill_title = PatternFill("solid", fgColor="1F4E79")  # dark blue
    fill_role = PatternFill("solid", fgColor="2F5597")  # blue
    fill_header = PatternFill("solid", fgColor="D9E1F2")  # light blue
    fill_category = PatternFill("solid", fgColor="F2F2F2")  # light gray

    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_role = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    font_body = Font(name="Calibri", size=11, color="1F1F1F")
    font_category = Font(name="Calibri", size=11, bold=True, color="1F1F1F")

    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_left_center = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Layout
    max_level = max(levels) if levels else 4
    level_cols = list(range(1, max_level + 1))
    cols = ["Nhóm năng lực"] + [f"Level {i}" for i in level_cols]
    last_col = len(cols)

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    t = ws.cell(1, 1, "P2 PERFORMANCE EVALUATION – PE Criteria Detail (theo Role)")
    t.fill = fill_title
    t.font = font_title
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Legend / hint
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    h = ws.cell(2, 1, "Gợi ý: dùng nút +/- (Outline) để thu gọn/mở rộng theo từng Role.")
    h.font = Font(name="Calibri", size=10, italic=True, color="404040")
    h.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    r = 4
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.applyStyles = True

    # Column widths
    widths = {1: 18}
    for i in range(2, last_col + 1):
        widths[i] = 48
    _set_col_widths(ws, widths)

    # Build per-role tables
    table_idx = 1
    category_label = {
        "Core_Skill": "A. Năng lực kỹ thuật",
        "Delivery": "B. Ownership & trách nhiệm",
        "DevOps": "C. DevOps & vận hành",
        "AI_Agent": "D. AI Agent & cải tiến",
    }
    for role in roles:
        # Role header row (summary row)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)
        c = ws.cell(r, 1, role)
        c.fill = fill_role
        c.font = font_role
        c.alignment = align_left_center
        ws.row_dimensions[r].height = 22

        role_header_row = r
        r += 1

        # Table header row
        for col_idx, label in enumerate(cols, start=1):
            cell_ = ws.cell(r, col_idx, label)
            cell_.fill = fill_header
            cell_.font = font_header
            cell_.alignment = align_center
        ws.row_dimensions[r].height = 20

        table_header_row = r
        r += 1

        # Category rows
        for cat in categories:
            cat_cell = ws.cell(r, 1, category_label.get(cat, cat.replace("_", " ")))
            cat_cell.fill = fill_category
            cat_cell.font = font_category
            cat_cell.alignment = align_left_center

            for i, lvl in enumerate(level_cols, start=2):
                desc = cell.get((role, cat, lvl), "")
                d = ws.cell(r, i, desc)
                d.font = font_body
                d.alignment = align_left_top
            ws.row_dimensions[r].height = 72
            r += 1

        table_first_row = table_header_row
        table_last_row = r - 1
        table_ref = f"A{table_first_row}:{get_column_letter(last_col)}{table_last_row}"

        # Add Excel table (filter + nicer styling)
        tab = Table(displayName=f"PE_{role}_Criteria_{table_idx}", ref=table_ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tab)
        table_idx += 1

        _apply_border(ws, f"A{table_header_row}:{get_column_letter(last_col)}{table_last_row}", border_thin)

        # Group rows under the role header row for overview
        detail_start = table_header_row
        detail_end = table_last_row
        for rr in range(detail_start, detail_end + 1):
            ws.row_dimensions[rr].outlineLevel = 1
            ws.row_dimensions[rr].hidden = True
        ws.row_dimensions[detail_end].collapsed = True
        ws.row_dimensions[role_header_row].outlineLevel = 0

        r += 1  # spacer

    # Freeze panes so title + tip stay visible
    ws.freeze_panes = "A4"

    wb.save(path)
    wb.close()


if __name__ == "__main__":
    main()

