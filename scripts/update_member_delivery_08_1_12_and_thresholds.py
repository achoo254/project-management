# -*- coding: utf-8 -*-
"""
Cập nhật file member (Bui Van Ho):
1. C25,C26,C27: Điểm Chưa đạt=0.8, Đạt=1, Tốt=1.2 (công thức B30).
2. C31: Ngưỡng xếp loại <1 Chưa đạt; 1 Đạt; >1 và ≤1.2 Tốt.
3. Sheet Cấu trúc & Trọng số: bảng ngưỡng chỉ 3 dòng (<1; 1; >1 và ≤1.2).
"""
from __future__ import annotations

import sys

import openpyxl

MEMBER_PATH = r"c:\CONG VIEC\project-management\members\P2_Form_BP_Cong_Nghe - Bui Van Ho.xlsx"
SHEET_PE = "Product_Engineer"
SHEET_CT = "Cấu trúc & Trọng số"

# P2-Delivery: Chưa đạt=0.8, Đạt=1, Tốt=1.2
DELIVERY_CHUA_DAT = 0.8
DELIVERY_DAT = 1
DELIVERY_TOT = 1.2


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    wb = openpyxl.load_workbook(MEMBER_PATH, data_only=False)

    # ----- Product_Engineer -----
    if SHEET_PE in wb.sheetnames:
        ws = wb[SHEET_PE]

        # 1. B30 – P2-Delivery: Chưa đạt=0.8, Đạt=1, Tốt=1.2
        part25 = f'IF(C25="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C25="Đạt",{DELIVERY_DAT},IF(C25="Tốt",{DELIVERY_TOT},0)))'
        part26 = f'IF(C26="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C26="Đạt",{DELIVERY_DAT},IF(C26="Tốt",{DELIVERY_TOT},0)))'
        part27 = f'IF(C27="Chưa đạt",{DELIVERY_CHUA_DAT},IF(C27="Đạt",{DELIVERY_DAT},IF(C27="Tốt",{DELIVERY_TOT},0)))'
        count_part = "(C25<>\"\")+(C26<>\"\")+(C27<>\"\")"
        ws.cell(30, 2).value = (
            f'=IF({count_part}=0,"",(({part25})+({part26})+({part27}))/({count_part}))'
        )

        # 1b. C24 – Mô tả thang điểm C25:C27: đổi thành 0.8, 1, 1.2
        if ws.cell(24, 3).value and "1.5" in str(ws.cell(24, 3).value):
            ws.cell(24, 3).value = "Điểm (Chưa đạt=0.8, Đạt=1, Tốt=1.2)"

        # 2. C31 – Ngưỡng: <1 Chưa đạt; 1 Đạt; >1 và ≤1.2 Tốt
        for m in list(ws.merged_cells.ranges):
            if m.min_row <= 31 <= m.max_row and m.min_col <= 3 <= m.max_col:
                ws.unmerge_cells(str(m))
                break
        ws.cell(31, 3).value = '=IF(B31="","",IF(B31<1,"Chưa đạt",IF(B31<=1,"Đạt",IF(B31<=1.2,"Tốt",""))))'

    # ----- Cấu trúc & Trọng số: bảng ngưỡng 3 dòng -----
    if SHEET_CT in wb.sheetnames:
        ws_ct = wb[SHEET_CT]
        # Tìm dòng "NGƯỠNG XẾP LOẠI" hoặc "P2+ Tổng" trong cột A
        header_row = None
        for r in range(1, 50):
            val = ws_ct.cell(r, 1).value
            if val and "NGƯỠNG XẾP LOẠI" in str(val).upper():
                header_row = r + 1  # bảng bắt đầu ngay dưới
                break
            if val and str(val).strip() == "P2+ Tổng" and ws_ct.cell(r, 2).value and "Xếp loại" in str(ws_ct.cell(r, 2).value):
                header_row = r
                break
        if header_row is None:
            for r in range(1, 50):
                if ws_ct.cell(r, 1).value and str(ws_ct.cell(r, 1).value).strip() == "P2+ Tổng":
                    header_row = r
                    break
        if header_row is not None:
            # Ghi đè: dòng 1 = header, 2-4 = 3 mốc
            ws_ct.cell(header_row, 1).value = "P2+ Tổng"
            ws_ct.cell(header_row, 2).value = "Xếp loại"
            ws_ct.cell(header_row + 1, 1).value = "< 1"
            ws_ct.cell(header_row + 1, 2).value = "Chưa đạt"
            ws_ct.cell(header_row + 2, 1).value = "1"
            ws_ct.cell(header_row + 2, 2).value = "Đạt"
            ws_ct.cell(header_row + 3, 1).value = "> 1 và ≤ 1.2"
            ws_ct.cell(header_row + 3, 2).value = "Tốt"
            # Xóa dòng cũ thứ 5 (≥ 4.2 Xuất sắc) nếu có
            for c in (1, 2):
                if ws_ct.cell(header_row + 4, c).value is not None:
                    ws_ct.cell(header_row + 4, c).value = None

    wb.save(MEMBER_PATH)
    wb.close()
    print("Đã cập nhật: C25:C27 điểm 0.8/1/1.2, C31 ngưỡng <1/1/>1-1.2, sheet Cấu trúc & Trọng số 3 dòng ngưỡng.")


if __name__ == "__main__":
    main()
