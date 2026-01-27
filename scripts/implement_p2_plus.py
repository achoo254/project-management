"""
Script to implement P2+ model in Excel file.
P2+ merges OKR/P3 into P2 evaluation.
"""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _safe_str(v: object) -> str:
    return "" if v is None else str(v).strip()


def _apply_border(ws, cell_range: str, border: Border) -> None:
    """Apply border to a range of cells."""
    for row in ws[cell_range]:
        for c in row:
            c.border = border


def _set_col_widths(ws, widths: dict[int, float]) -> None:
    """Set column widths."""
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def create_executive_summary(wb: openpyxl.Workbook) -> None:
    """Create Executive Summary sheet with P2+ definition."""
    # Remove if exists
    if "Executive Summary" in wb.sheetnames:
        wb.remove(wb["Executive Summary"])
    
    # Create as first sheet
    ws = wb.create_sheet("Executive Summary", 0)
    
    # Styles
    fill_title = PatternFill("solid", fgColor="1F4E79")  # dark blue
    fill_section = PatternFill("solid", fgColor="2F5597")  # blue
    fill_body = PatternFill("solid", fgColor="FFFFFF")  # white
    
    font_title = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_body = Font(name="Calibri", size=11, color="1F1F1F")
    font_highlight = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    title = ws.cell(1, 1, "P2+ - MÔ HÌNH ĐÁNH GIÁ NĂNG LỰC MỞ RỘNG")
    title.fill = fill_title
    title.font = font_title
    title.alignment = align_center
    ws.row_dimensions[1].height = 40
    
    # Content
    row = 3
    content = """P2+ là mô hình đánh giá năng lực mở rộng dành riêng cho Bộ phận Công Nghệ iNET, trong đó kết quả thực hiện OKR/KPI được tích hợp như một thành phần của P2, thay vì tách thành P3 độc lập. Mục tiêu là đơn giản hóa quy trình, phản ánh đúng đặc thù công việc kỹ thuật và đảm bảo đánh giá cân bằng giữa năng lực dài hạn và kết quả ngắn hạn."""
    
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    cell = ws.cell(row, 1, content)
    cell.fill = fill_body
    cell.font = font_body
    cell.alignment = align_left_top
    ws.row_dimensions[row].height = 80
    
    row += 2
    
    # Important note
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    note_title = ws.cell(row, 1, "👉 Lưu ý quan trọng:")
    note_title.fill = fill_section
    note_title.font = font_section
    note_title.alignment = align_left_top
    ws.row_dimensions[row].height = 25
    
    row += 1
    
    note_content = """• Đây không phải "bỏ P3"
• Mà là biến thể có chủ đích, tích hợp OKR/KPI vào P2
• Giúp đánh giá toàn diện cả năng lực dài hạn (P2-Core) và kết quả ngắn hạn (P2-Delivery)"""
    
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    note_cell = ws.cell(row, 1, note_content)
    note_cell.fill = fill_body
    note_cell.font = font_body
    note_cell.alignment = align_left_top
    ws.row_dimensions[row].height = 60
    
    # Column widths
    _set_col_widths(ws, {1: 20, 2: 20, 3: 20, 4: 20, 5: 20})
    
    # Freeze panes
    ws.freeze_panes = "A2"


def create_structure_sheet(wb: openpyxl.Workbook) -> None:
    """Create 'Cấu trúc & Trọng số' sheet."""
    # Remove if exists
    if "Cấu trúc & Trọng số" in wb.sheetnames:
        wb.remove(wb["Cấu trúc & Trọng số"])
    
    # Find position after Executive Summary
    exec_idx = wb.sheetnames.index("Executive Summary") if "Executive Summary" in wb.sheetnames else 0
    ws = wb.create_sheet("Cấu trúc & Trọng số", exec_idx + 1)
    
    # Styles
    fill_title = PatternFill("solid", fgColor="1F4E79")  # dark blue
    fill_section = PatternFill("solid", fgColor="2F5597")  # blue
    fill_header = PatternFill("solid", fgColor="D9E1F2")  # light blue
    fill_body = PatternFill("solid", fgColor="FFFFFF")  # white
    fill_category = PatternFill("solid", fgColor="F2F2F2")  # light gray
    
    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    font_body = Font(name="Calibri", size=11, color="1F1F1F")
    font_category = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    font_formula = Font(name="Calibri", size=12, bold=True, color="1F1F1F")
    
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_left_center = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    row = 1
    
    # Title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    title = ws.cell(row, 1, "CẤU TRÚC & TRỌNG SỐ P2+")
    title.fill = fill_title
    title.font = font_title
    title.alignment = align_center
    ws.row_dimensions[row].height = 35
    row += 2
    
    # Section 1: Cấu trúc tổng thể
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    section1 = ws.cell(row, 1, "1️⃣ CẤU TRÚC TỔNG THỂ P2+")
    section1.fill = fill_section
    section1.font = font_section
    section1.alignment = align_left_center
    ws.row_dimensions[row].height = 25
    row += 1
    
    # P2+ structure table
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    header = ws.cell(row, 1, "P2+ gồm 2 thành phần:")
    header.fill = fill_header
    header.font = font_header
    header.alignment = align_left_center
    ws.row_dimensions[row].height = 20
    row += 1
    
    # A. P2-Core
    ws.cell(row, 1, "A. P2-Core – Năng lực & cách làm việc (70%)").fill = fill_category
    ws.cell(row, 1).font = font_category
    ws.cell(row, 1).alignment = align_left_center
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 20
    row += 1
    
    core_items = [
        "A. Năng lực kỹ thuật",
        "B. Ownership & Trách nhiệm",
        "C. DevOps & Vận hành",
        "D. AI Agent & Cải tiến"
    ]
    
    for item in core_items:
        ws.cell(row, 2, f"  • {item}").font = font_body
        ws.cell(row, 2).alignment = align_left_center
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 18
        row += 1
    
    row += 1
    
    # B. P2-Delivery
    ws.cell(row, 1, "B. P2-Delivery – Kết quả công việc (30%)").fill = fill_category
    ws.cell(row, 1).font = font_category
    ws.cell(row, 1).alignment = align_left_center
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 20
    row += 1
    
    delivery_items = [
        "OKR/KPI tháng hoặc quý",
        "Chất lượng delivery"
    ]
    
    for item in delivery_items:
        ws.cell(row, 2, f"  • {item}").font = font_body
        ws.cell(row, 2).alignment = align_left_center
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 18
        row += 1
    
    row += 2
    
    # Section III: CÁCH TÍNH ĐIỂM P2+
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    section2 = ws.cell(row, 1, "III. CÁCH TÍNH ĐIỂM P2+")
    section2.fill = fill_section
    section2.font = font_section
    section2.alignment = align_left_center
    ws.row_dimensions[row].height = 25
    row += 1
    
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    intro_p2plus = ws.cell(row, 1, "P2+ là điểm đánh giá cuối cùng dùng để xếp loại năng lực.")
    intro_p2plus.fill = fill_body
    intro_p2plus.font = font_body
    intro_p2plus.alignment = align_left_center
    ws.row_dimensions[row].height = 22
    row += 1
    
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    formula = ws.cell(row, 1, "P2+ Tổng = (P2-Core × 70%) + (P2-Delivery × 30%)")
    formula.fill = fill_body
    formula.font = font_formula
    formula.alignment = align_center
    ws.row_dimensions[row].height = 30
    row += 1
    
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    note = ws.cell(row, 1, "📌 Trong đó:")
    note.font = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    note.alignment = align_left_center
    ws.row_dimensions[row].height = 20
    row += 1
    
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    note_core = ws.cell(row, 1, "P2-Core: phản ánh năng lực & cách làm việc dài hạn")
    note_core.font = font_body
    note_core.alignment = align_left_center
    ws.row_dimensions[row].height = 18
    row += 1
    
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    note_del = ws.cell(row, 1, "P2-Delivery: phản ánh kết quả thực tế trong kỳ")
    note_del.font = font_body
    note_del.alignment = align_left_center
    ws.row_dimensions[row].height = 18
    row += 2
    
    # Section IV: NGUYÊN TẮC ĐÁNH GIÁ (BẮT BUỘC TUÂN THỦ)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    section3 = ws.cell(row, 1, "IV. NGUYÊN TẮC ĐÁNH GIÁ (BẮT BUỘC TUÂN THỦ)")
    section3.fill = fill_section
    section3.font = font_section
    section3.alignment = align_left_center
    ws.row_dimensions[row].height = 25
    row += 1
    
    rules = [
        ("1️⃣ Không chấm trùng nội dung", [
            "P2-Core: chấm cách làm việc, năng lực, hành vi",
            "P2-Delivery: chấm kết quả đầu ra",
            "Ví dụ: Chủ động xử lý sự cố → P2-Core | Số lượng sự cố giảm → P2-Delivery"
        ]),
        ("2️⃣ OKR không thay thế năng lực", [
            "Nhân sự có kết quả OKR cao nhưng vi phạm chuẩn kỹ thuật, gây rủi ro hệ thống hoặc ảnh hưởng đội nhóm sẽ không được xếp loại P2+ cao."
        ]),
        ("3️⃣ Đánh giá theo mặt bằng chung", [
            "Điểm số phản ánh mức độ ổn định và phổ biến",
            "Không đánh giá theo 1–2 tình huống nổi bật"
        ])
    ]
    
    for rule_title, rule_items in rules:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        rule_header = ws.cell(row, 1, rule_title)
        rule_header.fill = fill_header
        rule_header.font = font_header
        rule_header.alignment = align_left_center
        ws.row_dimensions[row].height = 20
        row += 1
        
        for item in rule_items:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            rule_cell = ws.cell(row, 2, f"  • {item}")
            rule_cell.font = font_body
            rule_cell.alignment = align_left_top
            ws.row_dimensions[row].height = 20
            row += 1
        
        row += 1
    
    row += 1
    
    # Section V: Ngưỡng xếp loại
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    section4 = ws.cell(row, 1, "V. NGƯỠNG XẾP LOẠI")
    section4.fill = fill_section
    section4.font = font_section
    section4.alignment = align_left_center
    ws.row_dimensions[row].height = 25
    row += 1
    
    # Threshold table
    thresholds = [
        ("P2+ Tổng", "Xếp loại"),
        ("< 2.8", "Chưa đạt"),
        ("2.8 – 3.4", "Đạt"),
        ("3.5 – 4.1", "Tốt"),
        ("≥ 4.2", "Xuất sắc")
    ]
    
    for i, (col1, col2) in enumerate(thresholds):
        cell1 = ws.cell(row, 1, col1)
        cell2 = ws.cell(row, 2, col2)
        
        if i == 0:  # Header
            cell1.fill = fill_header
            cell2.fill = fill_header
            cell1.font = font_header
            cell2.font = font_header
        else:
            cell1.fill = fill_body
            cell2.fill = fill_body
            cell1.font = font_body
            cell2.font = font_body
        
        cell1.alignment = align_center
        cell2.alignment = align_center
        ws.row_dimensions[row].height = 20
        row += 1
    
    # Apply borders to threshold table
    _apply_border(ws, f"A{row-len(thresholds)}:B{row-1}", border_thin)
    
    # Column widths
    _set_col_widths(ws, {1: 25, 2: 30, 3: 30, 4: 30})
    
    # Freeze panes
    ws.freeze_panes = "A2"


def find_product_engineer_structure(ws) -> dict:
    """Find current structure of Product_Engineer sheet."""
    structure = {
        "last_row": ws.max_row,
        "block_d_end_row": None,
        "block_d_start_row": None,
        "p2_core_formula_cell": None,
        "p2_core_label_row": None,
        "p2_final_cell": None,
        "insertion_row": None
    }
    
    block_d_start = None
    block_d_end = None
    
    # Find Block D
    for row in range(1, ws.max_row + 1):
        cell_value = _safe_str(ws.cell(row, 1).value)
        
        # Look for Block D header
        if ("D." in cell_value and "AI Agent" in cell_value) or "D. AI Agent &" in cell_value:
            block_d_start = row
            structure["block_d_start_row"] = row
            
            # Find the end of Block D - look for empty rows or next section
            for r in range(row + 1, min(row + 100, ws.max_row + 1)):
                next_val = _safe_str(ws.cell(r, 1).value)
                # Check if we hit P2 calculations or next block
                if "P2" in next_val and ("Core" in next_val or "Final" in next_val or "Tổng" in next_val):
                    block_d_end = r - 1
                    break
                # Check for empty section (multiple empty rows)
                if not next_val and r > row + 5:
                    # Check next few rows
                    empty_count = 0
                    for check_r in range(r, min(r + 5, ws.max_row + 1)):
                        if not _safe_str(ws.cell(check_r, 1).value):
                            empty_count += 1
                    if empty_count >= 2:
                        block_d_end = r - 1
                        break
                # Check for next block
                if "E." in next_val:
                    block_d_end = r - 1
                    break
            
            if block_d_end is None:
                # If we didn't find a clear end, look for P2 calculations
                for r in range(row + 10, min(row + 50, ws.max_row + 1)):
                    next_val = _safe_str(ws.cell(r, 1).value)
                    if "P2" in next_val:
                        block_d_end = r - 2
                        break
                
                if block_d_end is None:
                    block_d_end = min(row + 30, ws.max_row)
            
            structure["block_d_end_row"] = block_d_end
            break
    
    # Find P2 Core formula/label
    for row in range(1, ws.max_row + 1):
        cell_value = _safe_str(ws.cell(row, 1).value)
        
        # Find P2 Core label
        if "P2" in cell_value and "Core" in cell_value and "Delivery" not in cell_value:
            structure["p2_core_label_row"] = row
            # Check if column B has formula or value
            if ws.cell(row, 2).data_type == "f":  # formula
                structure["p2_core_formula_cell"] = f"{get_column_letter(2)}{row}"
            elif ws.cell(row, 2).value is not None:
                structure["p2_core_formula_cell"] = f"{get_column_letter(2)}{row}"
        
        # Find P2 Final (old P2 calculation)
        if "P2" in cell_value and ("Final" in cell_value or ("Tổng" in cell_value and "P2+" not in cell_value)):
            if ws.cell(row, 2).data_type == "f":
                structure["p2_final_cell"] = f"{get_column_letter(2)}{row}"
            elif ws.cell(row, 2).value is not None:
                structure["p2_final_cell"] = f"{get_column_letter(2)}{row}"
    
    # Determine insertion row (after Block D, before P2 calculations)
    if block_d_end:
        structure["insertion_row"] = block_d_end + 3
    elif structure["p2_core_label_row"]:
        structure["insertion_row"] = structure["p2_core_label_row"] - 5
    else:
        structure["insertion_row"] = ws.max_row - 10
    
    return structure


def add_block_e_to_product_engineer(wb: openpyxl.Workbook) -> None:
    """Add Block E (P2-Delivery) to Product_Engineer sheet."""
    if "Product_Engineer" not in wb.sheetnames:
        print("Warning: Product_Engineer sheet not found")
        return
    
    ws = wb["Product_Engineer"]
    structure = find_product_engineer_structure(ws)
    
    # Styles (matching existing blocks)
    fill_category = PatternFill("solid", fgColor="F2F2F2")  # light gray
    fill_header = PatternFill("solid", fgColor="D9E1F2")  # light blue
    fill_body = PatternFill("solid", fgColor="FFFFFF")  # white
    
    font_category = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    font_body = Font(name="Calibri", size=11, color="1F1F1F")
    
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    align_left_center = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Find insertion point (after Block D)
    if structure.get("insertion_row"):
        start_row = structure["insertion_row"]
    elif structure.get("block_d_end_row"):
        start_row = structure["block_d_end_row"] + 3
    else:
        start_row = ws.max_row - 10
    
    # Block E header
    ws.cell(start_row, 1, "E. Kết quả công việc (P2-Delivery – từ OKR/KPI)").fill = fill_category
    ws.cell(start_row, 1).font = font_category
    ws.cell(start_row, 1).alignment = align_left_center
    ws.row_dimensions[start_row].height = 25
    start_row += 1
    
    # Block E intro: phản ánh kết quả đầu ra, không đánh giá lại năng lực/cách làm việc
    intro = "Phần này phản ánh kết quả đầu ra, không đánh giá lại năng lực hay cách làm việc. (P2-Core là phần đánh giá NĂNG LỰC & CÁCH LÀM VIỆC mang tính nền tảng và dài hạn của nhân sự Công Nghệ.)"
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=3)
    ws.cell(start_row, 1, intro).font = font_body
    ws.cell(start_row, 1).alignment = align_left_top
    ws.row_dimensions[start_row].height = 42
    start_row += 1
    
    # Table header
    headers = ["Tiêu chí", "Mô tả", "Điểm (1–5)"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, col_idx, header)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
    ws.row_dimensions[start_row].height = 20
    start_row += 1
    
    # Criteria rows
    criteria = [
        ("Mức độ hoàn thành OKR", "Tỷ lệ hoàn thành các mục tiêu đã cam kết trong kỳ đánh giá"),
        ("Chất lượng kết quả", "Ít lỗi nghiêm trọng; hạn chế làm lại; không gây rủi ro hệ thống"),
        ("Cam kết & tiến độ", "Hoàn thành đúng hạn; chủ động báo sớm rủi ro; phối hợp tốt"),
    ]
    
    criteria_start_row = start_row
    for criterion, description in criteria:
        ws.cell(start_row, 1, criterion).font = font_body
        ws.cell(start_row, 1).alignment = align_left_center
        ws.cell(start_row, 2, description).font = font_body
        ws.cell(start_row, 2).alignment = align_left_center
        # Score cell (empty for input)
        score_cell = ws.cell(start_row, 3, "")
        score_cell.font = font_body
        score_cell.alignment = align_center
        ws.row_dimensions[start_row].height = 25
        start_row += 1
    
    criteria_end_row = start_row - 1
    
    # Apply borders
    _apply_border(ws, f"A{criteria_start_row-1}:C{criteria_end_row}", border_thin)
    
    # P2-Delivery formula note
    start_row += 1
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=3)
    ws.cell(start_row, 1, "👉 P2-Delivery = Trung bình cộng các tiêu chí trên").font = font_body
    ws.cell(start_row, 1).alignment = align_left_center
    ws.row_dimensions[start_row].height = 20
    start_row += 1
    
    # P2-Delivery calculation
    ws.cell(start_row, 1, "P2-Delivery").font = font_category
    ws.cell(start_row, 1).alignment = align_left_center
    
    # Formula: AVERAGE of the 3 criteria scores
    score_col = get_column_letter(3)
    formula = f"=AVERAGE({score_col}{criteria_start_row}:{score_col}{criteria_end_row})"
    ws.cell(start_row, 2, formula).font = font_body
    ws.cell(start_row, 2).alignment = align_center
    ws.cell(start_row, 2).number_format = "0.00"
    ws.row_dimensions[start_row].height = 25
    
    # Store P2-Delivery cell reference for later use
    p2_delivery_cell = f"{get_column_letter(2)}{start_row}"
    
    # Update P2+ Tổng formula
    start_row += 2
    
    # Find P2-Core value cell
    p2_core_ref = structure.get("p2_core_formula_cell")
    if p2_core_ref is None:
        # Try to find it by searching
        p2_core_row = structure.get("p2_core_label_row")
        if p2_core_row:
            p2_core_ref = f"{get_column_letter(2)}{p2_core_row}"
        else:
            # Search for P2 Core
            for r in range(1, ws.max_row + 1):
                val = _safe_str(ws.cell(r, 1).value)
                if "P2" in val and "Core" in val and "Delivery" not in val:
                    p2_core_ref = f"{get_column_letter(2)}{r}"
                    break
    
    # Add P2+ Tổng
    ws.cell(start_row, 1, "P2+ Tổng").font = font_category
    ws.cell(start_row, 1).alignment = align_left_center
    
    if p2_core_ref:
        # Formula: (P2-Core × 0.7) + (P2-Delivery × 0.3)
        p2_plus_formula = f"=({p2_core_ref}*0.7)+({p2_delivery_cell}*0.3)"
        ws.cell(start_row, 2, p2_plus_formula).font = font_body
        ws.cell(start_row, 2).alignment = align_center
        ws.cell(start_row, 2).number_format = "0.00"
    else:
        ws.cell(start_row, 2, "=Cần link đến P2-Core").font = font_body
    
    ws.row_dimensions[start_row].height = 25
    
    # Column widths
    if ws.column_dimensions["A"].width < 30:
        ws.column_dimensions["A"].width = 30
    if ws.column_dimensions["B"].width < 40:
        ws.column_dimensions["B"].width = 40
    if ws.column_dimensions["C"].width < 15:
        ws.column_dimensions["C"].width = 15


def update_summary_ranking(wb: openpyxl.Workbook) -> None:
    """Update Summary_Ranking sheet with new columns."""
    if "Summary_Ranking" not in wb.sheetnames:
        print("Warning: Summary_Ranking sheet not found")
        return
    
    ws = wb["Summary_Ranking"]
    
    # Find header row and data start
    header_row = None
    for row in range(1, min(30, ws.max_row + 1)):
        cell_val = _safe_str(ws.cell(row, 1).value)
        # Check multiple possible header patterns
        if any(keyword in cell_val for keyword in ["Nhân sự", "Tên", "Họ tên", "Họ và tên", "Tên nhân viên", "Employee", "Name"]):
            header_row = row
            break
        # Also check if row has multiple filled cells (likely header)
        filled_count = sum(1 for col in range(1, min(10, ws.max_column + 1)) if ws.cell(row, col).value is not None)
        if filled_count >= 3 and row <= 5:
            header_row = row
            break
    
    if header_row is None:
        # Try row 1 as default
        if ws.cell(1, 1).value is not None:
            header_row = 1
        else:
            print("Warning: Could not find header row in Summary_Ranking, skipping update")
            return
    
    # Find last column
    last_col = 1
    for col in range(1, ws.max_column + 1):
        if ws.cell(header_row, col).value is not None:
            last_col = col
    
    # Styles
    fill_header = PatternFill("solid", fgColor="D9E1F2")  # light blue
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Add new headers
    new_headers = ["P2-Core", "OKR (P2-Delivery)", "P2+ Tổng"]
    for i, header in enumerate(new_headers, start=1):
        col = last_col + i
        cell = ws.cell(header_row, col, header)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
    
    # Add formulas for each data row
    # Find Product_Engineer sheet structure to get cell references
    if "Product_Engineer" not in wb.sheetnames:
        return
    
    pe_ws = wb["Product_Engineer"]
    pe_structure = find_product_engineer_structure(pe_ws)
    
    # Find P2-Core, P2-Delivery, P2+ Tổng cells in Product_Engineer
    p2_core_ref = None
    p2_delivery_ref = None
    p2_plus_ref = None
    
    # Search for labels first
    for row in range(1, pe_ws.max_row + 1):
        val = _safe_str(pe_ws.cell(row, 1).value)
        col_b_letter = get_column_letter(2)
        
        if "P2-Core" in val or ("P2" in val and "Core" in val and "Delivery" not in val and "P2+" not in val):
            # Check if column B has a value or formula
            if pe_ws.cell(row, 2).value is not None or pe_ws.cell(row, 2).data_type == "f":
                p2_core_ref = f"Product_Engineer!{col_b_letter}{row}"
        
        if "P2-Delivery" in val:
            if pe_ws.cell(row, 2).value is not None or pe_ws.cell(row, 2).data_type == "f":
                p2_delivery_ref = f"Product_Engineer!{col_b_letter}{row}"
        
        if "P2+ Tổng" in val or ("P2+" in val and "Tổng" in val):
            if pe_ws.cell(row, 2).value is not None or pe_ws.cell(row, 2).data_type == "f":
                p2_plus_ref = f"Product_Engineer!{col_b_letter}{row}"
    
    # If not found by label, search by formula pattern
    if not p2_core_ref:
        # Look for P2 calculation that doesn't include 0.7/0.3 (old P2 formula)
        for row in range(1, pe_ws.max_row + 1):
            if pe_ws.cell(row, 2).data_type == "f":
                formula = str(pe_ws.cell(row, 2).value)
                # Old P2 formula has percentages like 0.3, 0.25, 0.2 but not 0.7/0.3
                if ("0.3" in formula or "0.25" in formula or "0.2" in formula) and "0.7" not in formula:
                    val = _safe_str(pe_ws.cell(row, 1).value)
                    if "P2" in val or not val:  # Could be P2 or empty label
                        p2_core_ref = f"Product_Engineer!{get_column_letter(2)}{row}"
                        break
    
    if not p2_delivery_ref:
        # Look for AVERAGE formula near Block E
        for row in range(1, pe_ws.max_row + 1):
            if pe_ws.cell(row, 2).data_type == "f":
                formula = str(pe_ws.cell(row, 2).value)
                if "AVERAGE" in formula.upper():
                    # Check if nearby is Block E
                    for check_row in range(max(1, row-10), min(row+1, pe_ws.max_row + 1)):
                        check_val = _safe_str(pe_ws.cell(check_row, 1).value)
                        if "E." in check_val and "Kết quả" in check_val:
                            p2_delivery_ref = f"Product_Engineer!{get_column_letter(2)}{row}"
                            break
                    if p2_delivery_ref:
                        break
    
    if not p2_plus_ref:
        # Look for formula with 0.7 and 0.3
        for row in range(1, pe_ws.max_row + 1):
            if pe_ws.cell(row, 2).data_type == "f":
                formula = str(pe_ws.cell(row, 2).value)
                if "0.7" in formula and "0.3" in formula:
                    p2_plus_ref = f"Product_Engineer!{get_column_letter(2)}{row}"
                    break
    
    # Add formulas to data rows (assuming data starts at header_row + 1)
    data_start = header_row + 1
    for row in range(data_start, ws.max_row + 1):
        # Check if this is a data row (has name in first column)
        if ws.cell(row, 1).value is None:
            continue
        
        # Add formulas
        if p2_core_ref:
            ws.cell(row, last_col + 1, f"={p2_core_ref}").number_format = "0.00"
        if p2_delivery_ref:
            ws.cell(row, last_col + 2, f"={p2_delivery_ref}").number_format = "0.00"
        if p2_plus_ref:
            ws.cell(row, last_col + 3, f"={p2_plus_ref}").number_format = "0.00"
    
    # Adjust column widths
    for i in range(1, 4):
        col = last_col + i
        ws.column_dimensions[get_column_letter(col)].width = 15


def main() -> None:
    """Main function to implement P2+."""
    # Avoid UnicodeEncodeError on Windows consoles with legacy codepages.
    import sys
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    
    path = r"c:\CONG VIEC\project-management\P2_Form_Professional_Report_JuniorPlus_Table.xlsx"
    
    print("Loading workbook...")
    wb = openpyxl.load_workbook(path)
    
    print("Creating Executive Summary sheet...")
    create_executive_summary(wb)
    
    print("Creating 'Cấu trúc & Trọng số' sheet...")
    create_structure_sheet(wb)
    
    print("Adding Block E to Product_Engineer sheet...")
    add_block_e_to_product_engineer(wb)
    
    print("Updating Summary_Ranking sheet...")
    update_summary_ranking(wb)
    
    print("Saving workbook...")
    wb.save(path)
    wb.close()
    
    print("✅ P2+ implementation completed!")


if __name__ == "__main__":
    main()
