# -*- coding: utf-8 -*-
"""Inspect Product_Engineer sheet in P2_Form_BP_Cong_Nghe.xlsx"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import openpyxl
path = r"c:\CONG VIEC\project-management\P2_Form_BP_Cong_Nghe.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
print("Sheets:", wb.sheetnames)
if "Product_Engineer" in wb.sheetnames:
    ws = wb["Product_Engineer"]
    print("Max row:", ws.max_row, "Max col:", ws.max_column)
    for r in range(1, min(ws.max_row + 1, 150)):
        row_vals = []
        for c in range(1, min(ws.max_column + 1, 10)):
            v = ws.cell(r, c).value
            row_vals.append(str(v)[:40] if v is not None else "")
        if any(row_vals):
            print(r, "|", " | ".join(row_vals))
    # Data validations
    if ws.data_validations:
        for dv in ws.data_validations.dataValidation:
            print("DV:", dv.type, dv.formula1, getattr(dv, "sqref", ""))
    # Key formulas
    for r in [7, 8, 9, 10, 30]:
        for c in [2, 3, 4]:
            v = ws.cell(r, c).value
            if v and str(v).startswith("="):
                print(f"  Formula {get_column_letter(c)}{r}:", str(v)[:70])
wb.close()
