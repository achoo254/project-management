# -*- coding: utf-8 -*-
"""
Cập nhật sheet TBP_CÔNG NGHỆ trong P2_Form_BP_Cong_Nghe - LEAD.xlsx:
- Công thức Điểm quy đổi (D7:D10): từ 3 level (L1/L2/L3) lên 4 level (L1/L2/L3/L4).
- Thang điểm quy đổi: L1=1.5, L2=2.5, L3=3.5, L4=5.
- Data validation C7:C10: dropdown L1,L2,L3,L4.
- Dòng 13: THANG ĐÁNH GIÁ CHI TIẾT (L1/L2/L3/L4).
"""
from __future__ import annotations

import sys

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

LEAD_PATH = r"c:\CONG VIEC\project-management\P2_Form_BP_Cong_Nghe - LEAD.xlsx"
SHEET_TBP = "TBP_CÔNG NGHỆ"

# 4 level: L1=1.5, L2=2.5, L3=3.5, L4=5 (thang max 5)
FORMULA_CONVERT_4LEVEL = (
    'IF({cell}="L1",1.5,IF({cell}="L2",2.5,IF({cell}="L3",3.5,IF({cell}="L4",5,0))))'
)


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    wb = openpyxl.load_workbook(LEAD_PATH, data_only=False)
    if SHEET_TBP not in wb.sheetnames:
        print(f"Sheet '{SHEET_TBP}' not found.")
        wb.close()
        return

    ws = wb[SHEET_TBP]

    # 1. Công thức Điểm quy đổi D7:D10 (4 level)
    for r in range(7, 11):
        cell_c = f"C{r}"
        conv = FORMULA_CONVERT_4LEVEL.format(cell=cell_c)
        ws.cell(r, 4).value = f'=IF(C{r}="","",B{r}*{conv}/100)'

    # 2. Data validation C7:C10 -> L1,L2,L3,L4 (xóa bản cũ cho C7:C10 rồi thêm mới)
    to_remove = []
    for dv in list(ws.data_validations.dataValidation):
        sqref_str = str(dv.sqref) if dv.sqref else ""
        if "C7:C10" in sqref_str and "L1" in (dv.formula1 or ""):
            to_remove.append(dv)
    for dv in to_remove:
        ws.data_validations.dataValidation.remove(dv)
    dv_new = DataValidation(
        type="list",
        formula1='"L1,L2,L3,L4"',
        allow_blank=True,
    )
    dv_new.add("C7:C10")
    ws.add_data_validation(dv_new)

    # 3. Dòng 13: THANG ĐÁNH GIÁ CHI TIẾT (L1/L2/L3/L4)
    if ws.cell(13, 1).value and "L1/L2/L3" in str(ws.cell(13, 1).value):
        ws.cell(13, 1).value = "THANG ĐÁNH GIÁ CHI TIẾT (L1/L2/L3/L4)"

    wb.save(LEAD_PATH)
    wb.close()
    print("Đã cập nhật sheet TBP_CÔNG NGHỆ: công thức 4 level (L1=1.5, L2=2.5, L3=3.5, L4=5), dropdown L1/L2/L3/L4, tiêu đề thang L1/L2/L3/L4.")


if __name__ == "__main__":
    main()
