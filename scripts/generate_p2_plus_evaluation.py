# -*- coding: utf-8 -*-
"""
Generate P2+ Evaluation Excel workbook for Technology Department.
All content in Vietnamese. Output: P2_Plus_Evaluation_Tech_Department.xlsx
"""
from __future__ import annotations

import os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule

# Sheet names (Vietnamese)
SH_GUIDE = "Hướng dẫn"
SH_CRITERIA = "Tiêu chí theo nhóm"
SH_MEMBERS = "Thành viên"
SH_SELF = "Tự đánh giá"
SH_KPI = "KPI Công việc"
SH_FINAL = "Điểm tổng"
SH_DASH = "Tổng quan"

MAX_DATA_ROW = 51  # rows 2..51 = 50 members


def _apply_border(ws, cell_range: str, border: Border) -> None:
    for row in ws[cell_range]:
        for c in row:
            c.border = border


def _set_col_widths(ws, widths: dict[int, float]) -> None:
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def create_sheet_guide(wb: openpyxl.Workbook) -> None:
    """Tạo sheet Hướng dẫn cách nhập điểm số."""
    ws = wb.create_sheet(SH_GUIDE, 0)
    fill_title = PatternFill("solid", fgColor="1F4E79")
    fill_section = PatternFill("solid", fgColor="2F5597")
    fill_body = PatternFill("solid", fgColor="FFFFFF")
    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    font_body_style = Font(name="Calibri", size=11, color="1F1F1F")
    align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    row = 1
    ws.cell(row, 1, "Hướng dẫn cách nhập điểm số - Đánh giá P2+")
    ws.cell(row, 1).fill = fill_title
    ws.cell(row, 1).font = font_title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.row_dimensions[row].height = 32
    row += 2

    sections = [
        ("1. Thứ tự nhập liệu",
         "• Bước 1: Nhập danh sách thành viên tại sheet \"Thành viên\" (Họ tên, Vai trò, Mức hiện tại, Mức mục tiêu).\n"
         "  - Vai trò: Backend, Frontend, Fullstack, Designer. Một người có thể có nhiều vai trò (nhập cách nhau bằng dấu phẩy, ví dụ: Backend, Frontend).\n"
         "• Bước 2: Nhập điểm Tự đánh giá tại sheet \"Tự đánh giá\" (điểm 1–3 cho từng nhóm).\n"
         "• Bước 3: Nhập điểm KPI tại sheet \"KPI Công việc\" (Đúng hạn, Chất lượng, Trách nhiệm).\n"
         "• Sheet \"Điểm tổng\" và \"Tổng quan\" tự động tính từ dữ liệu đã nhập."),
        ("2. Cách nhập điểm Tự đánh giá (sheet Tự đánh giá)",
         "• Mỗi hàng tương ứng một thành viên (Họ tên lấy tự động từ sheet Thành viên).\n"
         "• Bốn cột điểm: Kỹ thuật, Công việc, Trách nhiệm, AI Agent.\n"
         "• Chỉ nhập số nguyên từ 1 đến 3:\n"
         "  - 1 = Mức 1 (L1): Thực hiện độc lập có hướng dẫn\n"
         "  - 2 = Mức 2 (L2): Tối ưu, chủ động\n"
         "  - 3 = Mức 3 (L3): Thiết kế, dẫn dắt, tư duy hệ thống\n"
         "• Tham chiếu tiêu chí chi tiết tại sheet \"Tiêu chí theo nhóm\"."),
        ("3. Cách nhập điểm KPI Công việc (sheet KPI Công việc)",
         "• Ba cột cần nhập: Đúng hạn (40%), Chất lượng (30%), Trách nhiệm (30%).\n"
         "• Nhập điểm theo thang 1–5 (1 = thấp, 5 = cao) hoặc thang 1–3 tùy quy định nội bộ.\n"
         "• Cột \"Điểm Công việc\" tự tính: Đúng hạn x 0,4 + Chất lượng x 0,3 + Trách nhiệm x 0,3.\n"
         "• Không nhập chữ hoặc để trống nếu chưa đánh giá; có thể nhập 0 nếu cần."),
        ("4. Cách đọc kết quả",
         "• Sheet \"Điểm tổng\": Điểm Năng lực = trung bình 4 nhóm Tự đánh giá; Điểm P2+ Tổng = 50% Năng lực + 50% Delivery (Điểm Công việc).\n"
         "• Màu nền cột Điểm P2+ Tổng: Xanh (>= 2,5), Vàng (2,0 – 2,49), Đỏ (< 2,0).\n"
         "• Sheet \"Tổng quan\": Số lượng theo mức L1/L2/L3, điểm trung bình phòng ban, bảng Xuất sắc và Cần cải thiện."),
        ("5. Liên kết với file P2_Form_BP_Cong_Nghe.xlsx",
         "• Mối quan hệ: File này (P2_Plus_Evaluation_Tech_Department.xlsx) dùng để đánh giá P2+ nội bộ; file P2_Form_BP_Cong_Nghe.xlsx là mẫu đánh giá P2 Bộ phận Công Nghệ chính thức. Có thể dùng kết quả từ file này để điền hoặc báo cáo vào P2_Form.\n"
         "• Cách 1 – Tham chiếu từ P2_Form sang file này: Trong P2_Form_BP_Cong_Nghe.xlsx, tại ô cần lấy điểm, nhập công thức tham chiếu ngoài, ví dụ:\n"
         "  ='[P2_Plus_Evaluation_Tech_Department.xlsx]Điểm tổng'!D2\n"
         "  (Lấy Điểm P2+ Tổng của hàng 2. Thay D2 bằng D3, D4,... tương ứng từng người. Hai file nên đặt cùng thư mục.)\n"
         "• Cách 2 – Copy kết quả sang P2_Form: Copy cột Họ tên (A) và Điểm P2+ Tổng (D) từ sheet \"Điểm tổng\" của file này, dán vào vị trí tương ứng trong P2_Form_BP_Cong_Nghe.xlsx (Dán giá trị để tránh phụ thuộc liên kết).\n"
         "• Lưu ý: Khi mở file có công thức liên kết ngoài, Excel sẽ hỏi \"Cập nhật liên kết\". Chọn Có nếu file kia cùng thư mục và cần lấy dữ liệu mới nhất. Nếu đổi tên hoặc di chuyển file, cần sửa lại đường dẫn trong công thức."),
    ]
    for title, content in sections:
        ws.cell(row, 1, title)
        ws.cell(row, 1).fill = fill_section
        ws.cell(row, 1).font = font_section
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = 22
        row += 1
        ws.cell(row, 1, content)
        ws.cell(row, 1).fill = fill_body
        ws.cell(row, 1).font = font_body_style
        ws.cell(row, 1).alignment = align_left_top
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = max(40, content.count("\n") * 16 + 20)
        row += 2

    _set_col_widths(ws, {1: 18, 2: 55})
    ws.freeze_panes = "A2"


def create_sheet_criteria(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet(SH_CRITERIA, 1)
    fill_header = PatternFill("solid", fgColor="D9E1F2")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    font_body = Font(name="Calibri", size=11, color="1F1F1F")
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = ["Nhóm", "Mức 1", "Mức 2", "Mức 3"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
    ws.row_dimensions[1].height = 22

    rows_criteria = [
        ("Kỹ thuật",
         "Triển khai feature end-to-end có hướng dẫn; kiểm thử cơ bản.",
         "Thiết kế API/DB/UI độc lập; tối ưu hiệu năng/bảo mật; kiểm thử vững.",
         "Thiết kế kiến trúc; refactor hệ thống lớn; hướng dẫn người khác."),
        ("Công việc",
         "Hoàn thành đúng hạn; tỷ lệ lỗi thấp.",
         "Ước lượng tin cậy; chia nhỏ công việc; bản release ổn định.",
         "Chịu trách nhiệm delivery dự án; phối hợp liên đội; kết quả dự đoán được."),
        ("Trách nhiệm",
         "Chịu trách nhiệm với task được giao.",
         "Chủ động nhận diện rủi ro; đề xuất cải tiến.",
         "Thúc đẩy sáng kiến; gỡ bỏ blocker; chịu trách nhiệm toàn phần."),
        ("AI Agent",
         "Dùng công cụ AI hỗ trợ viết code.",
         "Thiết kế prompt/automation/workflow.",
         "Xây hệ thống AI; tối ưu năng suất đội."),
    ]
    for r, (group, l1, l2, l3) in enumerate(rows_criteria, start=2):
        ws.cell(r, 1, group).font = font_body
        ws.cell(r, 1).alignment = align_left
        ws.cell(r, 2, l1).font = font_body
        ws.cell(r, 2).alignment = align_left
        ws.cell(r, 3, l2).font = font_body
        ws.cell(r, 3).alignment = align_left
        ws.cell(r, 4, l3).font = font_body
        ws.cell(r, 4).alignment = align_left
        ws.row_dimensions[r].height = 28
    _apply_border(ws, f"A1:D{1+len(rows_criteria)}", border_thin)
    _set_col_widths(ws, {1: 18, 2: 38, 3: 38, 4: 38})
    ws.freeze_panes = "A2"


def create_sheet_members(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet(SH_MEMBERS, 2)
    fill_header = PatternFill("solid", fgColor="D9E1F2")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = ["Họ tên", "Vai trò", "Mức hiện tại", "Mức mục tiêu", "Ghi chú"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
    ws.row_dimensions[1].height = 22

    for r in range(2, MAX_DATA_ROW + 1):
        for c in range(1, 6):
            ws.cell(r, c).border = border_thin
    _set_col_widths(ws, {1: 20, 2: 28, 3: 14, 4: 14, 5: 30})

    # Vai trò: Backend, Frontend, Fullstack, Designer - không dùng dropdown để cho phép nhiều vai trò (nhập cách nhau bằng dấu phẩy)
    dv_role_hint = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=255, allow_blank=True)
    dv_role_hint.prompt = "Vai trò: Backend, Frontend, Fullstack, Designer. Có thể nhập nhiều vai trò, cách nhau bằng dấu phẩy (vd: Backend, Frontend)"
    dv_role_hint.promptTitle = "Vai trò"
    dv_role_hint.add(f"B2:B{MAX_DATA_ROW}")
    ws.add_data_validation(dv_role_hint)

    dv_level = DataValidation(type="list", formula1='"L1,L2,L3"', allow_blank=True)
    dv_level.add(f"C2:C{MAX_DATA_ROW}")
    dv_level.add(f"D2:D{MAX_DATA_ROW}")
    ws.add_data_validation(dv_level)
    ws.freeze_panes = "A2"


def create_sheet_self_assessment(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet(SH_SELF, 3)
    fill_header = PatternFill("solid", fgColor="D9E1F2")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = ["Họ tên", "Kỹ thuật", "Công việc", "Trách nhiệm", "AI Agent"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
    ws.row_dimensions[1].height = 22

    for r in range(2, MAX_DATA_ROW + 1):
        ws.cell(r, 1, f"='{SH_MEMBERS}'!A{r}")
        for col in range(2, 6):
            ws.cell(r, col).border = border_thin
    dv_score = DataValidation(
        type="whole", operator="between", formula1=1, formula2=3, allow_blank=True
    )
    dv_score.add(f"B2:E{MAX_DATA_ROW}")
    ws.add_data_validation(dv_score)
    _set_col_widths(ws, {1: 20, 2: 12, 3: 12, 4: 12, 5: 12})
    ws.freeze_panes = "A2"


def create_sheet_kpi(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet(SH_KPI, 4)
    fill_header = PatternFill("solid", fgColor="D9E1F2")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "Họ tên",
        "Đúng hạn (40%)",
        "Chất lượng (30%)",
        "Trách nhiệm (30%)",
        "Điểm Công việc",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
    ws.row_dimensions[1].height = 22

    for r in range(2, MAX_DATA_ROW + 1):
        ws.cell(r, 1, f"='{SH_MEMBERS}'!A{r}")
        # Chỉ tính khi có ít nhất một ô nhập (tránh lỗi khi để trống)
        ws.cell(r, 5, f"=IF(COUNT(B{r}:D{r})=0,\"\",B{r}*0.4+C{r}*0.3+D{r}*0.3)")
        ws.cell(r, 5).number_format = "0.00"
        for col in range(1, 6):
            ws.cell(r, col).border = border_thin
    _set_col_widths(ws, {1: 20, 2: 16, 3: 16, 4: 16, 5: 18})
    ws.freeze_panes = "A2"


def create_sheet_final(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet(SH_FINAL, 5)
    fill_header = PatternFill("solid", fgColor="D9E1F2")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = ["Họ tên", "Điểm Năng lực", "Điểm Công việc", "Điểm P2+ Tổng"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
    ws.row_dimensions[1].height = 22

    for r in range(2, MAX_DATA_ROW + 1):
        ws.cell(r, 1, f"='{SH_MEMBERS}'!A{r}")
        # Điểm Năng lực: tránh #DIV/0! khi chưa nhập điểm Tự đánh giá
        ws.cell(r, 2, f"=IF(COUNT('{SH_SELF}'!B{r}:E{r})=0,\"\",AVERAGE('{SH_SELF}'!B{r}:E{r}))")
        ws.cell(r, 3, f"='{SH_KPI}'!E{r}")
        # Điểm P2+ Tổng: 50% Năng lực + 50% Delivery; chỉ tính khi B và C là số (tránh #VALUE!)
        ws.cell(r, 4, f"=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r})),B{r}*0.5+C{r}*0.5,\"\")")
        for col in range(1, 5):
            if col > 1:
                ws.cell(r, col).number_format = "0.00"
            ws.cell(r, col).border = border_thin
    ws.column_dimensions["A"].width = 20
    for col in range(2, 5):
        ws.column_dimensions[get_column_letter(col)].width = 16

    fill_green = PatternFill("solid", fgColor="C6EFCE")
    fill_yellow = PatternFill("solid", fgColor="FFEB9C")
    fill_red = PatternFill("solid", fgColor="FFC7CE")
    # Chỉ tô màu khi D là số (tránh tô ô trống)
    ws.conditional_formatting.add(
        f"D2:D{MAX_DATA_ROW}",
        FormulaRule(formula=["AND(ISNUMBER(D2),D2>=2.5)"], fill=fill_green, stopIfTrue=True),
    )
    ws.conditional_formatting.add(
        f"D2:D{MAX_DATA_ROW}",
        FormulaRule(formula=["AND(ISNUMBER(D2),D2>=2,D2<2.5)"], fill=fill_yellow, stopIfTrue=True),
    )
    ws.conditional_formatting.add(
        f"D2:D{MAX_DATA_ROW}",
        FormulaRule(formula=["AND(ISNUMBER(D2),D2<2)"], fill=fill_red, stopIfTrue=True),
    )
    ws.freeze_panes = "A2"


def create_sheet_dashboard(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet(SH_DASH, 6)
    fill_title = PatternFill("solid", fgColor="1F4E79")
    fill_section = PatternFill("solid", fgColor="2F5597")
    fill_header = PatternFill("solid", fgColor="D9E1F2")
    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="1F1F1F")
    font_body = Font(name="Calibri", size=11, color="1F1F1F")
    thin = Side(style="thin", color="BFBFBF")
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    row = 1
    ws.cell(row, 1, "Tổng quan đánh giá P2+ - Bộ phận Công nghệ")
    ws.cell(row, 1).fill = fill_title
    ws.cell(row, 1).font = font_title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 28
    row += 2

    ws.cell(row, 1, "Số lượng theo mức")
    ws.cell(row, 1).fill = fill_section
    ws.cell(row, 1).font = font_section
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 1
    level_data_start = row
    for i, level in enumerate(["L1", "L2", "L3"], start=0):
        ws.cell(row + i, 1, level)
        ws.cell(row + i, 2, f"=COUNTIF('{SH_MEMBERS}'!C2:C{MAX_DATA_ROW},\"{level}\")")
        ws.cell(row + i, 1).font = font_body
        ws.cell(row + i, 2).font = font_body
    level_data_end = row + 2
    row = level_data_end + 2

    ws.cell(row, 1, "Điểm trung bình phòng ban")
    ws.cell(row, 1).fill = fill_section
    ws.cell(row, 1).font = font_section
    row += 1
    # Tránh #DIV/0! khi chưa có điểm nào
    ws.cell(row, 1, f"=IFERROR(IF(COUNT('{SH_FINAL}'!D2:D{MAX_DATA_ROW})=0,\"\",AVERAGE('{SH_FINAL}'!D2:D{MAX_DATA_ROW})),\"\")")
    ws.cell(row, 1).number_format = "0.00"
    ws.cell(row, 1).font = font_body
    avg_row = row
    row += 3

    chart = BarChart()
    chart.type = "col"
    chart.title = "Phân bố theo mức"
    chart.y_axis.title = "Số lượng"
    data = Reference(ws, min_col=2, min_row=level_data_start, max_row=level_data_end)
    cats = Reference(ws, min_col=1, min_row=level_data_start, max_row=level_data_end)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    chart.width = 12
    chart.height = 8
    ws.add_chart(chart, f"E1")

    row_top = row
    ws.cell(row, 1, "Xuất sắc (điểm P2+ >= 2.5)")
    ws.cell(row, 1).fill = fill_section
    ws.cell(row, 1).font = font_section
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 1
    ws.cell(row, 1, "Họ tên")
    ws.cell(row, 2, "Điểm P2+ Tổng")
    ws.cell(row, 1).fill = fill_header
    ws.cell(row, 2).fill = fill_header
    ws.cell(row, 1).font = font_header
    ws.cell(row, 2).font = font_header
    top_header_row = row
    row += 1
    top_start = row
    for r in range(2, MAX_DATA_ROW + 1):
        i = row - top_start
        r_ref = 2 + i
        # Tránh lỗi khi D trống hoặc lỗi: chỉ hiển thị khi D là số và >= 2.5
        ws.cell(row, 1, f"=IFERROR(IF(AND(ISNUMBER('{SH_FINAL}'!D{r_ref}),'{SH_FINAL}'!D{r_ref}>=2.5),'{SH_FINAL}'!A{r_ref},\"\"),\"\")")
        ws.cell(row, 2, f"=IFERROR(IF(AND(ISNUMBER('{SH_FINAL}'!D{r_ref}),'{SH_FINAL}'!D{r_ref}>=2.5),'{SH_FINAL}'!D{r_ref},\"\"),\"\")")
        ws.cell(row, 2).number_format = "0.00"
        row += 1
    top_end = row - 1
    row += 2

    ws.cell(row, 1, "Cần cải thiện (điểm P2+ < 2.0)")
    ws.cell(row, 1).fill = fill_section
    ws.cell(row, 1).font = font_section
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 1
    ws.cell(row, 1, "Họ tên")
    ws.cell(row, 2, "Điểm P2+ Tổng")
    ws.cell(row, 1).fill = fill_header
    ws.cell(row, 2).fill = fill_header
    ws.cell(row, 1).font = font_header
    ws.cell(row, 2).font = font_header
    row += 1
    need_start = row
    for r in range(2, MAX_DATA_ROW + 1):
        i = row - need_start
        r_ref = 2 + i
        # Tránh lỗi khi D trống hoặc lỗi: chỉ hiển thị khi D là số và < 2
        ws.cell(row, 1, f"=IFERROR(IF(AND(ISNUMBER('{SH_FINAL}'!D{r_ref}),'{SH_FINAL}'!D{r_ref}<2),'{SH_FINAL}'!A{r_ref},\"\"),\"\")")
        ws.cell(row, 2, f"=IFERROR(IF(AND(ISNUMBER('{SH_FINAL}'!D{r_ref}),'{SH_FINAL}'!D{r_ref}<2),'{SH_FINAL}'!D{r_ref},\"\"),\"\")")
        ws.cell(row, 2).number_format = "0.00"
        row += 1

    _set_col_widths(ws, {1: 28, 2: 18})
    ws.freeze_panes = "A2"


def main() -> None:
    try:
        import sys
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    create_sheet_guide(wb)
    create_sheet_criteria(wb)
    create_sheet_members(wb)
    create_sheet_self_assessment(wb)
    create_sheet_kpi(wb)
    create_sheet_final(wb)
    create_sheet_dashboard(wb)

    out_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out_path = os.path.join(out_dir, "P2_Plus_Evaluation_Tech_Department.xlsx")
    wb.save(out_path)
    wb.close()
    print(f"Đã tạo file: {out_path}")


if __name__ == "__main__":
    main()
